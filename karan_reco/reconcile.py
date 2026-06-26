import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

# ── Colours ───────────────────────────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
TITLE_FILL  = PatternFill("solid", fgColor="2E75B6")
OK_FILL     = PatternFill("solid", fgColor="E2EFDA")
WARN_FILL   = PatternFill("solid", fgColor="FFF2CC")
ERR_FILL    = PatternFill("solid", fgColor="FFE0E0")
ALT_FILL    = PatternFill("solid", fgColor="F5F9FF")
WHITE_FILL  = PatternFill("solid", fgColor="FFFFFF")

HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT  = Font(name="Calibri", bold=True, color="FFFFFF", size=13)
BOLD        = Font(name="Calibri", bold=True, size=10)
NORMAL      = Font(name="Calibri", size=10)

CENTER = Alignment(horizontal="center", vertical="center")
LEFT   = Alignment(horizontal="left",   vertical="center")
RIGHT  = Alignment(horizontal="right",  vertical="center")

def thin_border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)

def set_col_width(ws, col, width):
    ws.column_dimensions[get_column_letter(col)].width = width

def sc(cell, font=None, fill=None, align=None, border=None, num_fmt=None):
    if font:    cell.font            = font
    if fill:    cell.fill            = fill
    if align:   cell.alignment       = align
    if border:  cell.border          = border
    if num_fmt: cell.number_format   = num_fmt

# ══════════════════════════════════════════════════════════════════════════════
# 1. READ ORIGINAL DATA
# ══════════════════════════════════════════════════════════════════════════════

wb_src = openpyxl.load_workbook("E:/Saifullah/fabric-pos/karan_reco/STG STOCK & SALES.xlsx")
ws_sl  = wb_src["STG STOCKLIST"]

stocklist = []
orig_listed_sold = {}
for row in ws_sl.iter_rows(min_row=4, values_only=True):
    sno, item, pcs, orig, current, sold = (list(row) + [None]*6)[:6]
    if sno and item:
        item = item.strip()
        stocklist.append({"item": item, "pcs": pcs or 0, "original": orig or 0})
        orig_listed_sold[item] = sold if isinstance(sold, (int, float)) else None

wb_sales = openpyxl.load_workbook("E:/Saifullah/fabric-pos/karan_reco/STG SALES.xlsx")
ws_sales = wb_sales["Sheet1"]

sales_rows = []
last_date = last_customer = None
for row in ws_sales.iter_rows(min_row=3, values_only=True):
    sno, date, customer, item, bales, price, total = (list(row) + [None]*7)[:7]
    if item is None or bales is None:
        continue
    if date:     last_date     = date
    if customer: last_customer = customer
    cust = customer if customer else last_customer
    amt  = bales * price if isinstance(bales, (int,float)) and isinstance(price, (int,float)) else total
    sales_rows.append({
        "sno": sno, "date": last_date, "customer": cust,
        "item": item.strip(), "bales": bales, "price": price, "total": amt,
    })

ALIAS = {
    "Brand R/N T-Shirt":            "Brand Round Neck T-Shirt",
    "Brand Hoodie Sweatshirt":      "Brand Hoodie Sweatshirt",
    "Brand Hoodie Sweatshirt ":     "Brand Hoodie Sweatshirt",
    "Brand Round Neck Sweatshirt ": "Brand Round Neck Sweatshirt",
    "Brand Sweatshirt":             "Brand Round Neck Sweatshirt",
    "Brand Cap":                    "Brand Caps",
}

actual_sold = defaultdict(float)
for s in sales_rows:
    canonical = ALIAS.get(s["item"], s["item"])
    actual_sold[canonical] += s["bales"]

stocked_items = {r["item"] for r in stocklist}
extra_items = []
for item in actual_sold:
    if item not in stocked_items:
        extra_items.append({"item": item, "pcs": 0, "original": 0})

all_stock = stocklist + extra_items

# ══════════════════════════════════════════════════════════════════════════════
# 2. BUILD WORKBOOK
# ══════════════════════════════════════════════════════════════════════════════
wb = openpyxl.Workbook()

# ─────────────────────────────────────────────────────────────────────────────
# SHEET 1: CORRECTED STOCKLIST
# ─────────────────────────────────────────────────────────────────────────────
ws1 = wb.active
ws1.title = "STG STOCKLIST (Corrected)"
ws1.sheet_view.showGridLines = False
ws1.freeze_panes = "A3"

ws1.merge_cells("A1:H1")
t = ws1["A1"]
t.value = "STG STOCKLIST — RECONCILED"
sc(t, font=TITLE_FONT, fill=TITLE_FILL, align=CENTER)
ws1.row_dimensions[1].height = 28

headers = ["Sno","Item","Pcs/Bale","Original Stock (Bales)",
           "Actual Sold (Bales)","Actual Current Stock","Old Listed Sold","Status"]
ws1.row_dimensions[2].height = 30
for ci, h in enumerate(headers, 1):
    c = ws1.cell(row=2, column=ci, value=h)
    sc(c, font=HEADER_FONT, fill=HEADER_FILL, border=thin_border())
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for i, rec in enumerate(all_stock, 1):
    item     = rec["item"]
    pcs      = rec["pcs"]
    orig     = rec["original"]
    act_sold = actual_sold.get(item, 0)
    act_curr = orig - act_sold
    lst_sold = orig_listed_sold.get(item)
    is_new   = item not in stocked_items

    if is_new:
        status   = "MISSING FROM STOCKLIST"
        row_fill = ERR_FILL
    elif lst_sold is None and act_sold > 0:
        status   = "FORMULA — STOCK UPDATED"
        row_fill = WARN_FILL
    elif lst_sold is None:
        status   = "FORMULA — VERIFIED OK"
        row_fill = OK_FILL
    elif lst_sold != act_sold:
        status   = "MISMATCH (was %d)" % int(lst_sold)
        row_fill = ERR_FILL
    else:
        status   = "OK"
        row_fill = OK_FILL if i % 2 == 0 else ALT_FILL

    r = i + 2
    lst_disp = int(lst_sold) if lst_sold is not None else "formula"
    vals = [i, item, pcs, orig, act_sold, act_curr, lst_disp, status]
    for ci, val in enumerate(vals, 1):
        c = ws1.cell(row=r, column=ci, value=val)
        sc(c, font=NORMAL, fill=row_fill, border=thin_border())
        c.alignment = CENTER if ci in (1,3,4,5,6,7) else LEFT

tr = len(all_stock) + 3
for col in range(1, 9):
    c = ws1.cell(row=tr, column=col)
    c.fill = HEADER_FILL
    c.border = thin_border()
ws1.cell(row=tr, column=1, value="TOTAL")
sc(ws1.cell(row=tr, column=1), font=HEADER_FONT, align=CENTER)

total_orig = sum(r["original"] for r in all_stock)
total_sold = sum(actual_sold.get(r["item"], 0) for r in all_stock)
total_curr = total_orig - total_sold
for col, val in [(4, total_orig), (5, total_sold), (6, total_curr)]:
    c = ws1.cell(row=tr, column=col, value=val)
    sc(c, font=HEADER_FONT, fill=HEADER_FILL, align=CENTER, border=thin_border())

for col, w in [(1,5),(2,38),(3,10),(4,18),(5,18),(6,18),(7,14),(8,26)]:
    set_col_width(ws1, col, w)

# ─────────────────────────────────────────────────────────────────────────────
# SHEET 2: FULL SALES DETAIL
# ─────────────────────────────────────────────────────────────────────────────
ws2 = wb.create_sheet("STG SALES DETAIL")
ws2.sheet_view.showGridLines = False
ws2.freeze_panes = "A3"

ws2.merge_cells("A1:G1")
t2 = ws2["A1"]
t2.value = "STG SALES DETAIL — COMPLETE (Feb–Jun 2026)"
sc(t2, font=TITLE_FONT, fill=TITLE_FILL, align=CENTER)
ws2.row_dimensions[1].height = 28

hdrs2 = ["Sno","Date","Customer","Item","Bales","Price","Total Amount"]
ws2.row_dimensions[2].height = 22
for ci, h in enumerate(hdrs2, 1):
    c = ws2.cell(row=2, column=ci, value=h)
    sc(c, font=HEADER_FONT, fill=HEADER_FILL, align=CENTER, border=thin_border())

for i, s in enumerate(sales_rows, 1):
    r = i + 2
    row_fill = WHITE_FILL if i % 2 == 0 else ALT_FILL
    vals = [s["sno"], s["date"], s["customer"], s["item"],
            s["bales"], s["price"], s["total"]]
    for ci, val in enumerate(vals, 1):
        c = ws2.cell(row=r, column=ci, value=val)
        sc(c, font=NORMAL, fill=row_fill, border=thin_border())
        if ci == 2:
            c.number_format = "DD-MMM-YYYY"
            c.alignment = CENTER
        elif ci in (1, 5):
            c.alignment = CENTER
        elif ci in (6, 7):
            c.alignment = RIGHT
            c.number_format = "#,##0"
        else:
            c.alignment = LEFT

tr2 = len(sales_rows) + 3
total_bales  = sum(s["bales"] for s in sales_rows if isinstance(s["bales"], (int,float)))
total_amount = sum(s["total"] for s in sales_rows if isinstance(s["total"], (int,float)))
for col in range(1, 8):
    ws2.cell(row=tr2, column=col).fill   = HEADER_FILL
    ws2.cell(row=tr2, column=col).border = thin_border()
sc(ws2.cell(row=tr2, column=1, value="TOTAL"), font=HEADER_FONT, align=CENTER)
for col, val, fmt in [(5, total_bales, "#,##0"), (7, total_amount, "#,##0")]:
    c = ws2.cell(row=tr2, column=col, value=val)
    sc(c, font=HEADER_FONT, fill=HEADER_FILL, align=CENTER, border=thin_border())
    c.number_format = fmt

for col, w in [(1,5),(2,14),(3,16),(4,38),(5,8),(6,12),(7,15)]:
    set_col_width(ws2, col, w)

# ─────────────────────────────────────────────────────────────────────────────
# SHEET 3: DISCREPANCY REPORT
# ─────────────────────────────────────────────────────────────────────────────
ws3 = wb.create_sheet("Discrepancy Report")
ws3.sheet_view.showGridLines = False

ws3.merge_cells("A1:G1")
t3 = ws3["A1"]
t3.value = "DISCREPANCY REPORT — Generated 06-Jun-2026"
sc(t3, font=TITLE_FONT, fill=TITLE_FILL, align=CENTER)
ws3.row_dimensions[1].height = 28

hdrs3 = ["Item","Original Stock","Listed Sold","Actual Sold",
         "Listed Curr Stock","Actual Curr Stock","Issue"]
ws3.row_dimensions[2].height = 22
for ci, h in enumerate(hdrs3, 1):
    c = ws3.cell(row=2, column=ci, value=h)
    sc(c, font=HEADER_FONT, fill=HEADER_FILL, align=CENTER, border=thin_border())

issues = []
for rec in all_stock:
    item     = rec["item"]
    orig     = rec["original"]
    act_sold = actual_sold.get(item, 0)
    act_curr = orig - act_sold
    lst_sold = orig_listed_sold.get(item)
    is_new   = item not in stocked_items

    if is_new:
        issue = "MISSING FROM STOCKLIST — sold but never added to stocklist"
        fill  = ERR_FILL
    elif lst_sold is None and act_sold > 0:
        issue = "Broken formula in Sold cell — stock updated in corrected file"
        fill  = WARN_FILL
    elif isinstance(lst_sold, (int,float)) and lst_sold != act_sold:
        issue = "Sold count wrong: listed %d, actual %d" % (int(lst_sold), int(act_sold))
        fill  = ERR_FILL
    else:
        continue

    lst_curr = orig - (lst_sold or 0) if not is_new else None
    issues.append((item, orig if not is_new else "—",
                   int(lst_sold) if lst_sold is not None else "formula",
                   int(act_sold), lst_curr if lst_curr is not None else "—",
                   int(act_curr), issue, fill))

for i, row_data in enumerate(issues, 1):
    r = i + 2
    fill = row_data[7]
    for ci, val in enumerate(row_data[:7], 1):
        c = ws3.cell(row=r, column=ci, value=val)
        sc(c, font=NORMAL, fill=fill, border=thin_border())
        c.alignment = CENTER if ci in (2,3,4,5,6) else LEFT

for col, w in [(1,38),(2,14),(3,12),(4,12),(5,16),(6,16),(7,55)]:
    set_col_width(ws3, col, w)

# ── Save ──────────────────────────────────────────────────────────────────────
out = "E:/Saifullah/fabric-pos/karan_reco/STG RECONCILED.xlsx"
wb.save(out)
print("Saved:", out)
print("  Sheet 1 — Stocklist: %d items (%d newly added)" % (len(all_stock), len(extra_items)))
print("  Sheet 2 — Sales:     %d entries" % len(sales_rows))
print("  Sheet 3 — Issues:    %d discrepancies flagged" % len(issues))
print("  Total bales sold:    %d" % int(total_bales))
print("  Total sales amount:  %s" % "{:,}".format(int(total_amount)))
